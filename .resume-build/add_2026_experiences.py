from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/mustafa/paper")
INPUT = ROOT / "resume.docx"
OUTPUT = ROOT / "resume.docx"


def replace_paragraph_text_preserving_format(paragraph: Paragraph, text: str) -> None:
    text_nodes = paragraph._p.xpath(".//w:t")
    if not text_nodes:
        run = paragraph.add_run(text)
        run.font.name = "Times New Roman"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
        return

    text_nodes[0].text = text
    for node in text_nodes[1:]:
        node.text = ""


def populate_experience_table(table: Table, company: str, date: str, title: str) -> None:
    replace_paragraph_text_preserving_format(table.cell(0, 0).paragraphs[0], company)
    replace_paragraph_text_preserving_format(table.cell(0, 1).paragraphs[0], date)
    replace_paragraph_text_preserving_format(table.cell(1, 0).paragraphs[0], title)
    replace_paragraph_text_preserving_format(table.cell(1, 1).paragraphs[0], "")


doc = Document(INPUT)

cbre_table = next(
    table for table in doc.tables if table.cell(0, 0).text.strip() == "CBRE"
)
replace_paragraph_text_preserving_format(
    cbre_table.cell(1, 0).paragraphs[0], "UX Designer (Contract)"
)
cbre_bullet = next(
    paragraph
    for paragraph in doc.paragraphs
    if paragraph.text.startswith("Designed and built SmartFM's 3D login experience")
)
cbre_spacer = Paragraph(cbre_bullet._p.getnext(), cbre_bullet._parent)

entries = [
    (
        "Stealth Startup",
        "July 2026",
        "Design Engineer Consultant",
        [
            "Advised the founding team on product, design, and marketing strategy.",
            "Designed and built a low-latency frontend for an agentic trading platform for its public launch.",
        ],
    ),
    (
        "NextWork",
        "June 2026",
        "Design Advisor",
        [
            "Built a CSS and JavaScript motion library to explain technical concepts to learners.",
            "Designed microinteractions to make learning more enjoyable.",
        ],
    ),
]

existing_companies = {table.cell(0, 0).text.strip() for table in doc.tables}
for company, date, title, bullets in entries:
    if company not in existing_companies:
        table_xml = deepcopy(cbre_table._tbl)
        table = Table(table_xml, cbre_table._parent)
        populate_experience_table(table, company, date, title)
        cbre_table._tbl.addprevious(table_xml)

        for bullet in bullets:
            bullet_xml = deepcopy(cbre_bullet._p)
            bullet_paragraph = Paragraph(bullet_xml, cbre_bullet._parent)
            replace_paragraph_text_preserving_format(bullet_paragraph, bullet)
            cbre_table._tbl.addprevious(bullet_xml)

        cbre_table._tbl.addprevious(deepcopy(cbre_spacer._p))

# Keep the expanded resume to one page without changing its visual system.
section = doc.sections[0]
section.top_margin = Inches(0.50)
section.bottom_margin = Inches(0.50)

for style_name in ("Normal", "List Paragraph"):
    style = doc.styles[style_name]
    style.font.name = "Times New Roman"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.font.size = Pt(10.5)

for paragraph in doc.paragraphs:
    if paragraph.style.name == "List Paragraph" and paragraph.text.strip():
        paragraph.paragraph_format.line_spacing = 1.08
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
    elif paragraph.style.name == "Title":
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(3)
    elif not paragraph.text.strip():
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(4.5)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        if not paragraph.runs:
            paragraph.add_run("")
        for run in paragraph.runs:
            run.font.size = Pt(4.5)

for table in doc.tables:
    table.autofit = False
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.08
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)

for paragraph in doc.paragraphs:
    if paragraph.style.name == "Normal" and paragraph.text.strip():
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)

doc.save(OUTPUT)
