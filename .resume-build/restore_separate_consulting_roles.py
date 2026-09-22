from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.table import Table


RESUME = Path("/Users/mustafa/paper/resume.docx")


def find_table(document, heading: str):
    for table in document.tables:
        if table.cell(0, 0).text.strip() == heading:
            return table
    raise ValueError(f"Could not find table: {heading}")


def find_paragraph(document, text: str):
    for paragraph in document.paragraphs:
        if paragraph.text == text:
            return paragraph
    raise ValueError(f"Could not find paragraph: {text}")


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def set_experience_header(table, company: str, date: str) -> None:
    set_cell_text(table.cell(0, 0), company)
    set_cell_text(table.cell(0, 1), date)
    set_cell_text(table.cell(1, 0), "Design Engineer Consultant")
    set_cell_text(table.cell(1, 1), "Austin, TX")


def remove_paragraph(paragraph) -> None:
    element = paragraph._p
    element.getparent().remove(element)


document = Document(RESUME)
consulting_table = find_table(document, "Independent Consulting")
stealth_heading = find_paragraph(document, "Stealth Startup")
nextwork_heading = find_paragraph(document, "NextWork")
stealth_second_bullet = find_paragraph(
    document,
    "Designed and built a low-latency frontend for an agentic trading platform for its public launch.",
)
nextwork_second_bullet = find_paragraph(
    document,
    "Designed microinteractions to make learning more enjoyable.",
)

set_experience_header(consulting_table, "Stealth Startup", "July 2026")

# Clone the existing, fully formatted experience table so geometry and type
# remain identical to every other entry in the resume.
nextwork_table_xml = deepcopy(consulting_table._tbl)
nextwork_table = Table(nextwork_table_xml, consulting_table._parent)
set_experience_header(nextwork_table, "NextWork", "June 2026")

# Reuse the existing post-entry spacer, then insert the NextWork table between
# the two sets of bullets.
spacer_xml = deepcopy(nextwork_second_bullet._p.getnext())
stealth_second_bullet._p.addnext(spacer_xml)
spacer_xml.addnext(nextwork_table_xml)

remove_paragraph(stealth_heading)
remove_paragraph(nextwork_heading)
document.save(RESUME)

print(f"Restored separate Stealth Startup and NextWork entries in {RESUME}")
