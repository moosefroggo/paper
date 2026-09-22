from pathlib import Path

from docx import Document


RESUME_PATH = Path("/Users/mustafa/paper/resume.docx")
REPLACEMENTS = {
    (
        "Built a CSS and JavaScript motion library to explain technical concepts "
        "to learners."
    ): (
        "Built a motion library using the Canvas API and Three.js to explain "
        "technical concepts to learners."
    ),
    (
        "Built a native fleet safety and maintenance tool that saved more than "
        "$500K annually for Motive's largest US customer by fleet count."
    ): (
        "Built a fleet safety and maintenance product that saved $500K+ annually "
        "for the largest fleet operator in the U.S."
    ),
    (
        "Designed two enterprise products for developer onboarding and engineering "
        "team productivity, closing a $100K ARR contract."
    ): (
        "Designed developer onboarding and engineering productivity products that "
        "closed a $100K ARR contract."
    ),
}


def replace_text_preserving_paragraph_format(paragraph, new_text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


document = Document(RESUME_PATH)
updated = set()

for paragraph in document.paragraphs:
    text = paragraph.text.strip()
    for old_text, new_text in REPLACEMENTS.items():
        if text == old_text or text == new_text:
            replace_text_preserving_paragraph_format(paragraph, new_text)
            updated.add(new_text)
            break

missing = set(REPLACEMENTS.values()) - updated
if missing:
    raise RuntimeError(f"Could not find bullets to update: {sorted(missing)}")

document.save(RESUME_PATH)
