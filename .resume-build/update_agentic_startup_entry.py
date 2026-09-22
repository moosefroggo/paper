from pathlib import Path

from docx import Document


RESUME_PATH = Path("/Users/mustafa/paper/resume.docx")
OLD_HEADING = "Stealth Startup"
NEW_HEADING = "Agentic Trading Startup (Stealth)"
OLD_BULLET = (
    "Designed a human-agent trading workspace with Next.js, TypeScript, Zustand, "
    "and TanStack Query."
)
NEW_BULLET = (
    "Built a human-agent trading workspace with Next.js, TypeScript, Zustand, "
    "and TanStack Query."
)


def replace_text_preserving_paragraph_format(paragraph, new_text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


document = Document(RESUME_PATH)

heading_updated = False
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                if paragraph.text.strip() in {OLD_HEADING, NEW_HEADING}:
                    replace_text_preserving_paragraph_format(paragraph, NEW_HEADING)
                    heading_updated = True

bullet_updated = False
for paragraph in document.paragraphs:
    if paragraph.text.strip() in {OLD_BULLET, NEW_BULLET}:
        replace_text_preserving_paragraph_format(paragraph, NEW_BULLET)
        bullet_updated = True

if not heading_updated:
    raise RuntimeError("Could not find the stealth startup heading.")
if not bullet_updated:
    raise RuntimeError("Could not find the agentic trading bullet.")

document.save(RESUME_PATH)
