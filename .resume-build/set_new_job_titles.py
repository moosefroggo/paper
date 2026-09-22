from pathlib import Path

from docx import Document


resume_path = Path("/Users/mustafa/paper/resume.docx")
doc = Document(resume_path)

target_companies = {"Stealth Startup", "NextWork"}
for table in doc.tables:
    if table.cell(0, 0).text.strip() not in target_companies:
        continue

    paragraph = table.cell(1, 0).paragraphs[0]
    text_nodes = paragraph._p.xpath(".//w:t")
    text_nodes[0].text = "Design Engineer Consultant"
    for node in text_nodes[1:]:
        node.text = ""

    location_paragraph = table.cell(1, 1).paragraphs[0]
    location_nodes = location_paragraph._p.xpath(".//w:t")
    if location_nodes:
        location_nodes[0].text = "Austin, TX"
        for node in location_nodes[1:]:
            node.text = ""
    else:
        location_paragraph.add_run("Austin, TX")

doc.save(resume_path)
