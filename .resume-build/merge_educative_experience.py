from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path("/Users/mustafa/paper")
RESUME = ROOT / "resume.docx"


def replace_paragraph_text_preserving_format(paragraph: Paragraph, text: str) -> None:
    text_nodes = paragraph._p.xpath(".//w:t")
    if not text_nodes:
        run = paragraph.add_run(text)
        run.font.name = "Times New Roman"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
        return

    text_nodes[0].text = text
    for node in text_nodes[1:]:
        node.text = ""


doc = Document(RESUME)
educative_tables = [
    table for table in doc.tables if table.cell(0, 0).text.strip() == "Educative"
]

if len(educative_tables) == 2:
    lead_table, earlier_table = educative_tables
    uxrally_table = next(
        table for table in doc.tables if table.cell(0, 0).text.strip() == "UXrally"
    )

    replace_paragraph_text_preserving_format(
        lead_table.cell(0, 1).paragraphs[0], "September 2020 - September 2023"
    )
    replace_paragraph_text_preserving_format(
        lead_table.cell(1, 0).paragraphs[0], "Design Lead III"
    )
    replace_paragraph_text_preserving_format(
        lead_table.cell(1, 1).paragraphs[0], "Pakistan"
    )

    body = doc.element.body
    body_children = list(body)
    start = body_children.index(lead_table._tbl) + 1
    end = body_children.index(uxrally_table._tbl)
    existing_segment = body_children[start:end]

    bullet_template = next(
        Paragraph(node, lead_table._parent)
        for node in existing_segment
        if node.tag == qn("w:p")
        and Paragraph(node, lead_table._parent).style.name == "List Paragraph"
        and Paragraph(node, lead_table._parent).text.strip()
    )
    spacer_template = next(
        Paragraph(node, lead_table._parent)
        for node in reversed(existing_segment)
        if node.tag == qn("w:p")
        and not Paragraph(node, lead_table._parent).text.strip()
    )

    for node in existing_segment:
        body.remove(node)

    bullets = [
        "Joined Educative as the sole designer for 30 developers. Hired and led a team of 10 designers after Educative's $12M Series A.",
        "Designed core consumer learning products as Educative grew from 500K to 2.5M users.",
        "Designed two enterprise products for developer onboarding and engineering team productivity, closing a $100K ARR contract.",
        "Built and implemented the design system across 6 SaaS products.",
    ]

    for text in bullets:
        bullet_xml = deepcopy(bullet_template._p)
        bullet = Paragraph(bullet_xml, lead_table._parent)
        replace_paragraph_text_preserving_format(bullet, text)
        uxrally_table._tbl.addprevious(bullet_xml)

    uxrally_table._tbl.addprevious(deepcopy(spacer_template._p))

# Use the freed space to improve readability while preserving one page.
section = doc.sections[0]
section.top_margin = Inches(0.50)
section.bottom_margin = Inches(0.50)

for style_name in ("Normal", "List Paragraph"):
    style = doc.styles[style_name]
    style.font.size = Pt(11)

for paragraph in doc.paragraphs:
    if paragraph.style.name == "List Paragraph" and paragraph.text.strip():
        paragraph.paragraph_format.line_spacing = 1.12

for table in doc.tables:
    table.autofit = False
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.12

doc.save(RESUME)
