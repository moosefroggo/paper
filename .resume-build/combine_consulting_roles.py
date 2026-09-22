from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


RESUME = Path("/Users/mustafa/paper/resume.docx")


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def find_table(document, company: str):
    for table in document.tables:
        if table.cell(0, 0).text.strip() == company:
            return table
    raise ValueError(f"Could not find table for {company}")


def find_paragraph(document, text: str):
    for paragraph in document.paragraphs:
        if paragraph.text == text:
            return paragraph
    raise ValueError(f"Could not find paragraph: {text}")


def client_heading_before(paragraph, text: str) -> None:
    new_p = OxmlElement("w:p")
    p_pr = OxmlElement("w:pPr")

    style = OxmlElement("w:pStyle")
    style.set(qn("w:val"), "Normal")
    p_pr.append(style)

    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:line"), "269")
    spacing.set(qn("w:lineRule"), "auto")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    p_pr.append(spacing)
    new_p.append(p_pr)

    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    bold = OxmlElement("w:b")
    r_pr.append(bold)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "22")
    r_pr.append(size)
    run.append(r_pr)

    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    new_p.append(run)

    paragraph._p.addprevious(new_p)


def remove_block(block) -> None:
    element = block._element
    element.getparent().remove(element)


document = Document(RESUME)
stealth_table = find_table(document, "Stealth Startup")
nextwork_table = find_table(document, "NextWork")

# Retain the resume's existing two-row experience-header pattern while making
# the two short projects read as one consulting period.
set_cell_text(stealth_table.cell(0, 0), "Independent Consulting")
set_cell_text(stealth_table.cell(0, 1), "June 2026 - July 2026")
set_cell_text(stealth_table.cell(1, 0), "Design Engineer Consultant")
set_cell_text(stealth_table.cell(1, 1), "Austin, TX")

stealth_bullet = find_paragraph(
    document,
    "Advised the founding team on product, design, and marketing strategy.",
)
nextwork_bullet = find_paragraph(
    document,
    "Built a CSS and JavaScript motion library to explain technical concepts to learners.",
)

client_heading_before(stealth_bullet, "Stealth Startup")
client_heading_before(nextwork_bullet, "NextWork")

# Find the spacer directly after the second Stealth bullet robustly.
stealth_second = find_paragraph(
    document,
    "Designed and built a low-latency frontend for an agentic trading platform for its public launch.",
)
candidate = stealth_second._p.getnext()
if candidate is not None and candidate.tag == qn("w:p"):
    candidate_paragraph = Paragraph(candidate, document)
    if candidate_paragraph.text == " ":
        remove_block(candidate_paragraph)

remove_block(nextwork_table)
document.save(RESUME)

print(f"Updated {RESUME}")
