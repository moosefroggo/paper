from pathlib import Path

from docx import Document


SOURCE = Path("/Users/mustafa/paper/resume.docx")
OUTPUT = Path("/Users/mustafa/paper/resume-germany.docx")


REPLACEMENTS = {
    "Austin, TX": "Austin, TX · Moving to Germany Sep 1st",
    "PROFESSIONAL EXPERIENCE": "WORK EXPERIENCE",
    "August 2024 - April 2026": "08/2024 – 04/2026",
    "September 2013 - June 2017": "09/2013 – 06/2017",
    "June 2025 - August 2025": "06/2025 – 08/2025",
    "September 2023 - August 2024": "09/2023 – 08/2024",
    "May 2021 - September 2023": "05/2021 – 09/2023",
    "September 2020 - May 2021": "09/2020 – 05/2021",
    "December 2018 - September 2020": "12/2018 – 09/2020",
    "Master's in Human-Computer Interaction": "M.Sc. Human–Computer Interaction",
    "Bachelor's in Software Engineering": "B.Sc. Software Engineering",
}


def all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def replace_in_runs(paragraph):
    for run in paragraph.runs:
        for old, new in REPLACEMENTS.items():
            if old in run.text:
                run.text = run.text.replace(old, new)


def move_education_after_experience(document):
    body = document.element.body
    children = list(body)

    education_heading = next(
        element
        for element in children
        if "EDUCATION" in "".join(element.itertext())
    )
    work_heading = next(
        element
        for element in children
        if "WORK EXPERIENCE" in "".join(element.itertext())
    )
    skills_heading = next(
        element
        for element in children
        if "SKILLS" in "".join(element.itertext())
    )

    start = children.index(education_heading)
    end = children.index(work_heading)
    education_block = children[start:end]

    for element in education_block:
        body.remove(element)

    insertion_index = list(body).index(skills_heading)
    for offset, element in enumerate(education_block):
        body.insert(insertion_index + offset, element)


def main():
    document = Document(SOURCE)

    for paragraph in all_paragraphs(document):
        replace_in_runs(paragraph)

    move_education_after_experience(document)

    document.core_properties.title = "Mustafa Ali Akbar – Resume – Germany"
    document.core_properties.subject = "Product design resume for roles in Germany"
    document.core_properties.author = "Mustafa Ali Akbar"
    document.core_properties.keywords = "Product Design, UX Design, Germany"

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
