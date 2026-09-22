from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

base = Path('/Users/mustafa/paper')
path = base / 'Mustafa Ali Akbar - Chancenkarte CV.docx'
d = Document(path)
original = [c.text for t in d.tables for row in list(t.rows)[1 if t != d.tables[-1] else 0:] for c in row.cells]
p = d.paragraphs[0]
p.text = 'Mustafa Ali Akbar'
for r in p.runs:
    r.font.size = Pt(19)
    r.font.bold = True
p.paragraph_format.space_after = Pt(5)
d.styles['Title'].font.size = Pt(19)
for root in (d.element, d.styles.element):
    for el in list(root.iter(qn('w:pBdr'))):
        el.getparent().remove(el)

for ti, t in enumerate(d.tables):
    if ti < 3 and t.cell(0,0).text in ('Period', 'Project'):
        t._tbl.remove(t.rows[0]._tr)
    for el in list(t._tbl.iter(qn('w:shd'))):
        el.getparent().remove(el)
    for tag in ('w:tblBorders', 'w:tcBorders'):
        for el in list(t._tbl.iter(qn(tag))):
            el.getparent().remove(el)
    borders = OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'),'nil')
        borders.append(e)
    t._tbl.tblPr.append(borders)
    for row in t.rows:
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            pr = cell._tc.get_or_add_tcPr()
            for el in list(pr.findall(qn('w:tcMar'))):
                pr.remove(el)
            mar = OxmlElement('w:tcMar')
            for side, value in [('top',70),('bottom',150 if ti < 2 else 100),('left',0 if ci == 0 else 90),('right',50)]:
                e=OxmlElement('w:'+side)
                e.set(qn('w:w'),str(value)); e.set(qn('w:type'),'dxa')
                mar.append(e)
            pr.append(mar)
            for p in cell.paragraphs:
                for r in p.runs:
                    if r.font.color.rgb == RGBColor(255,255,255):
                        r.font.color.rgb = RGBColor(0,0,0)
for p in d.paragraphs:
    if p.style.name == 'Heading 1':
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
d.save(path)
print(path)
